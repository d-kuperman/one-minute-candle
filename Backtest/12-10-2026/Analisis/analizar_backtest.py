from pathlib import Path
import hashlib, json, warnings
import numpy as np
import pandas as pd
import openpyxl
from reportlab.graphics.shapes import Drawing, String, Line
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics import renderSVG
from reportlab.lib import colors
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
OUT=Path(__file__).parent
SRC=OUT.parent/'ReportTester-26316626.xlsx'
sha=hashlib.sha256(SRC.read_bytes()).hexdigest()
w=openpyxl.load_workbook(SRC,data_only=True); s=w['Sheet1']
orders={int(s.cell(r,2).value):[c.value for c in s[r]] for r in range(72,266)}
rows=[]; events=[]; active=None; balance=100000.; max_balance_error=0
for r in range(269,463):
    v=[c.value for c in s[r]]
    t=pd.to_datetime(v[0],format='%Y.%m.%d %H:%M:%S')
    oid=int(v[7]); order=orders[oid]
    assert order[11]=='filled' and order[2]==v[2]=='EURUSD'
    assert order[3]==v[3] and pd.to_datetime(order[9])==t
    assert float(order[4].split('/')[0])==float(v[5])==float(order[4].split('/')[1])
    balance+=sum(float(v[i]) for i in [8,9,10]); max_balance_error=max(max_balance_error,abs(balance-v[11]))
    events.append(dict(fecha=t,deal=int(v[1]),balance=v[11],flujo=sum(v[i] for i in [8,9,10])))
    if v[4]=='in':
        assert active is None
        active=(r,v,order,t)
    elif v[4]=='out':
        assert active is not None
        rr,a,o,ta=active; sign=1 if a[3]=='buy' else -1
        assert a[3]!=v[3] and a[5]==v[5] and t>=ta
        lot=float(a[5]); entry=float(a[6]); close=float(v[6]); sl=float(o[7])
        risk=round(sign*(entry-sl)*lot*100000,8)
        gross=round(a[10]+v[10],2); commission=round(a[8]+v[8],2); swap=round(a[9]+v[9],2)
        net=round(gross+commission+swap,2)
        assert risk>0 and abs(round(sign*(close-entry)*lot*100000,2)-gross)<.011
        assert v[12].startswith('sl ')
        endsl=float(v[12].split()[1]); moved=sign*(endsl-sl)>1e-9
        rows.append(dict(id=len(rows)+1,orden_entrada=int(a[7]),orden_salida=oid,deal_entrada=int(a[1]),deal_salida=int(v[1]),fila_entrada=rr,fila_salida=r,apertura=ta,cierre=t,direccion=a[3].upper(),lotes=lot,precio_entrada=entry,precio_salida=close,sl_orden_entrada=sl,sl_comentario_salida=endsl,sl_mejorado=moved,distancia_sl_pips=round(sign*(entry-sl)*10000,6),riesgo_precio_usd=risk,riesgo_con_comision_usd=risk-commission,bruto_usd=gross,comision_usd=commission,swap_usd=swap,neto_usd=net,resultado_r=net/risk,resultado_r_con_comision=net/(risk-commission),duracion_segundos=(t-ta).total_seconds(),desviacion_sl_pips=round(sign*(close-endsl)*10000,6)))
        active=None
    else: raise ValueError(v)
assert active is None and len(rows)==97 and len(orders)==194
assert len(set(e['deal'] for e in events))==194 and max_balance_error<.0001
d=pd.DataFrame(rows); e=pd.DataFrame(events)
d['acumulado_usd']=d.neto_usd.cumsum(); d['acumulado_r']=d.resultado_r.cumsum()
d['resultado']=np.select([d.neto_usd>0,d.neto_usd<0],['Ganadora','Perdedora'],default='Breakeven')
d['hora']=d.apertura.dt.hour; d['dia']=d.cierre.dt.strftime('%Y-%m-%d')
def metrics(g):
    x=g.neto_usd; pos=x[x>0]; neg=x[x<0]
    return dict(trades=len(g),ganadoras=len(pos),perdedoras=len(neg),be=int((x==0).sum()),winrate=100*len(pos)/len(g) if len(g) else None,neto=x.sum(),pf=pos.sum()/-neg.sum() if len(neg) else None,media=x.mean(),ganancia_media=pos.mean(),perdida_media=-neg.mean(),r=g.resultado_r.sum(),r_media=g.resultado_r.mean(),mejor=x.max(),peor=x.min())
def dd(x,start=0):
    a=np.r_[start,np.asarray(x)]; peak=np.maximum.accumulate(a); dif=peak-a
    return float(max(dif))
def streaks(x):
    groups=[]; sign=None; vals=[]
    for val in x:
        sig=np.sign(val)
        if sig!=sign:
            if vals: groups.append((int(sign),len(vals),round(sum(vals),2)))
            vals=[]; sign=sig
        vals.append(val)
    if vals: groups.append((int(sign),len(vals),round(sum(vals),2)))
    return groups
daily=[]
for day in pd.date_range('2026-10-12','2026-10-16'):
    key=str(day.date()); g=d[d.dia==key]; ev=e[e.fecha.dt.date==day.date()]
    start=100000. if not len(e[e.fecha<day]) else e[e.fecha<day].balance.iloc[-1]
    m=metrics(g); m.update(dia=key,dd=dd(ev.balance,start),dd_pct=max((np.maximum.accumulate(np.r_[start,ev.balance])-np.r_[start,ev.balance])/np.maximum.accumulate(np.r_[start,ev.balance]))*100)
    daily.append(m)
daily=pd.DataFrame(daily)
hour=pd.DataFrame([dict(hora=h,**metrics(d[d.hora==h])) for h in range(24)])
side=pd.DataFrame([dict(direccion=k,**metrics(g)) for k,g in d.groupby('direccion')])
d['grupo_duracion']=pd.cut(d.duracion_segundos,[-1,60,180,300,600,1200,np.inf],right=False,labels=['<1 min','[1,3) min','[3,5) min','[5,10) min','[10,20) min','≥20 min'])
duration=pd.DataFrame([dict(grupo=str(k),**metrics(g)) for k,g in d.groupby('grupo_duracion',observed=False)])
labels=['≤−1R','(−1,−0.05)R','[−0.05,+0.05]R','(+0.05,+1)R','[+1,+2)R','[+2,+3]R','>+3R']
r=d.resultado_r.round(8).to_numpy(); masks=[r<=-1,(r>-1)&(r<-.05),(r>=-.05)&(r<=.05),(r>.05)&(r<1),(r>=1)&(r<2),(r>=2)&(r<=3),r>3]
assert np.all(np.sum(masks,axis=0)==1)
rd=pd.DataFrame([dict(intervalo=l,trades=int(m.sum()),neto_usd=d.loc[m,'neto_usd'].sum(),r=d.loc[m,'resultado_r'].sum()) for l,m in zip(labels,masks)])
for name,df in [('Operaciones_Analizadas',d),('Resultados_Diarios',daily),('Resultados_Horarios',hour),('Resultados_Direccion',side),('Resultados_Duracion',duration),('Distribucion_R',rd),('Balance_Por_Deal',e)]: df.to_csv(OUT/(name+'.csv'),index=False,encoding='utf-8-sig',float_format='%.8f')
pos=d[d.neto_usd>0].sort_values('neto_usd',ascending=False)
peaks=np.maximum.accumulate(np.r_[100000,e.balance]); bals=np.r_[100000,e.balance]
grosspos=d.loc[d.bruto_usd>0,'bruto_usd'].sum(); grossneg=d.loc[d.bruto_usd<0,'bruto_usd'].sum()
summ=dict(source_sha256=sha,metrics=metrics(d),gross_positive=grosspos,gross_negative=grossneg,gross=d.bruto_usd.sum(),commission=d.comision_usd.sum(),swap=d.swap_usd.sum(),net_positive=pos.neto_usd.sum(),net_negative=d.loc[d.neto_usd<0,'neto_usd'].sum(),dd_balance=max(peaks-bals),dd_pct=max((peaks-bals)/peaks)*100,dd_absolute=100000-min(bals),dd_r=dd(d.acumulado_r),dd_closed=dd(100000+d.acumulado_usd,100000),streaks=streaks(d.neto_usd),streaks_gross=streaks(d.bruto_usd+d.swap_usd),top3=pos.head(3).neto_usd.sum(),top5=pos.head(5).neto_usd.sum(),duration=d.duracion_segundos.describe().to_dict(),lot=d.lotes.describe().to_dict(),risk=d.riesgo_precio_usd.describe().to_dict(),stop=d.distancia_sl_pips.describe().to_dict(),moved=int(d.sl_mejorado.sum()),unchanged=int((~d.sl_mejorado).sum()),near_minus1=int((abs(r+1)<=.05).sum()),near0=int((abs(r)<=.05).sum()),over={str(i):int((r>i).sum()) for i in [1,2,3,5]},duration_spearman=dict(statistic=d.duracion_segundos.rank().corr(d.neto_usd.rank())),mean_ci_iid=[d.neto_usd.mean()-1.985*d.neto_usd.std()/len(d)**.5,d.neto_usd.mean()+1.985*d.neto_usd.std()/len(d)**.5],balance_error=max_balance_error)
(OUT/'metricas_verificacion.json').write_text(json.dumps(summ,indent=2,ensure_ascii=False,default=lambda o:float(o)),encoding='utf-8')
def chart(name,title,subtitle,xlabel,ylabel,values,labels=None,line=False):
    drawing=Drawing(1000,510)
    drawing.add(String(75,480,title,fontName='Helvetica-Bold',fontSize=19,fillColor=colors.HexColor('#21344b')))
    drawing.add(String(75,457,subtitle,fontSize=11,fillColor=colors.HexColor('#566576')))
    drawing.add(String(75,432,ylabel,fontSize=11))
    c=LinePlot() if line else VerticalBarChart()
    c.x=80;c.y=95;c.width=875;c.height=315
    if line:
        c.data=[list(enumerate(map(float,values)))];c.lines[0].strokeColor=colors.HexColor('#2c5d91');c.lines[0].strokeWidth=2
        c.xValueAxis.valueMin=0;c.xValueAxis.valueMax=len(values)-1;c.xValueAxis.valueStep=max(1,round(len(values)/8))
        axis=c.yValueAxis
    else:
        c.data=[list(map(float,values))];c.categoryAxis.categoryNames=labels
        c.categoryAxis.labels.fontSize=9;c.categoryAxis.labels.angle=0
        c.bars[0].fillColor=colors.HexColor('#2c5d91')
        for i,val in enumerate(values): c.bars[(0,i)].fillColor=colors.HexColor('#17836b' if val>=0 else '#b4424d')
        axis=c.valueAxis
    lo=min(0,min(values));hi=max(0,max(values));span=hi-lo or 1
    axis.valueMin=lo-span*.06 if lo<0 else 0;axis.valueMax=hi+span*.1
    axis.labels.fontSize=10;axis.visibleGrid=1;axis.gridStrokeColor=colors.HexColor('#e2e6eb')
    drawing.add(c);drawing.add(String(520,35,xlabel,textAnchor='middle',fontSize=11))
    renderSVG.drawToFile(drawing,str(OUT/(name+'.svg')))
chart('01_Beneficio_Acumulado','Beneficio acumulado por ejecución','Incluye comisión de apertura y swap. No es una curva de equity.','Evento contable (0 = depósito inicial)','Beneficio acumulado (USD)',np.r_[0,e.balance-100000],line=True)
chart('02_Drawdown','Drawdown de balance','Máximo: 8,799.86 USD. Incluye débitos al abrir posiciones.','Evento contable (0 = depósito inicial)','Caída respecto del máximo previo (USD)',bals-peaks,line=True)
chart('03_Resultados_Diarios','Resultado neto por fecha de cierre','Viernes: sin ejecuciones; no se acredita cobertura del día.','Día de octubre de 2026','Resultado neto (USD)',daily.neto,['Lun 12','Mar 13','Mié 14','Jue 15','Vie 16*'])
chart('04_Resultados_Horarios','Resultado por hora de entrada','Hora original del XLSX. Zona horaria no identificada.','Hora de entrada (0-23)','Resultado neto (USD)',hour.neto,[str(h) for h in range(24)])
chart('05_BUY_SELL','Resultado neto por dirección','BUY: 41 operaciones. SELL: 56 operaciones.','Dirección de entrada','Resultado neto (USD)',side.neto,list(side.direccion))
hist,edges=np.histogram(d.neto_usd,bins=np.arange(-2500,10001,1000))
chart('06_Histograma_Resultados','Distribución de resultados netos','Intervalos de 1,000 USD. Etiquetas = punto medio del intervalo.','Resultado individual (USD), punto medio de cada intervalo','Operaciones (cantidad)',hist,[str(int(x)) for x in (edges[:-1]+edges[1:])/2])
chart('07_Distribucion_R','Distribución de resultados en R','Neto / riesgo inicial de precio individual. Cercano a cero: [-0.05,+0.05]R.','Intervalos exactos definidos en la sección 5 del informe','Operaciones (cantidad)',rd.trades,['<=-1','(-1,-.05)','[-.05,.05]','(.05,1)','[1,2)','[2,3]','>3'])
chart('08_Duracion','Duración de las posiciones','Precisión de los datos: un segundo. Intervalos sin solapamiento.','Intervalo de duración (minutos)','Operaciones (cantidad)',duration.trades,['<1','[1,3)','[3,5)','[5,10)','[10,20)','>=20'])
chart('09_Concentracion','Concentración de las ganancias netas','Top 3: 29.65%. Top 5: 37.63% de la suma de ganancias netas.','Índice de ganadora, de mayor a menor (0 = mayor ganadora)','Contribución acumulada (%)',100*pos.neto_usd.cumsum()/pos.neto_usd.sum(),line=True)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==sha
print('Validación completada: 97 posiciones, 194 deals, 194 órdenes. CSV y SVG generados.')
